"""Rabbani formatting: retain the source package and its relationships."""
from io import BytesIO
from copy import deepcopy
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from lxml import etree

ALIASES = {
 'ABSTRACT':'abstract','ABSTRAK':'abstrak',
 'PENDAHULUAN':'intro','INTRODUCTION':'intro',
 'METODE':'method','METODE PENELITIAN':'method','RESEARCH METHODOLOGY':'method','METHODS':'method','METHODOLOGY':'method',
 'HASIL DAN PEMBAHASAN':'results','HASIL PENELITIAN DAN PEMBAHASAN':'results','RESULT AND DISCUSSION':'results','RESULTS AND DISCUSSION':'results',
 'HASIL PENELITIAN':'results','PEMBAHASAN':'results',
 'KESIMPULAN':'conclusion','SIMPULAN':'conclusion','KESIMPULAN DAN SARAN':'conclusion','CONCLUSION':'conclusion',
 'DAFTAR PUSTAKA':'references','BIBLIOGRAPHY':'references','REFERENCES':'references'}

def heading(text):
    text=re.sub(r'^\s*(?:[A-Z]|\d+|[IVX]+)[.)]\s*','',text.strip(),flags=re.I)
    return ALIASES.get(text.upper().rstrip(':').strip())

def style(p,size=12,bold=None,spacing=1.5,align=Align.JUSTIFY,indent=0):
    f=p.paragraph_format
    f.keep_with_next=False;f.keep_together=False;f.page_break_before=False
    f.widow_control=True
    f.line_spacing=spacing;f.space_before=Pt(0);f.space_after=Pt(0)
    f.first_line_indent=Cm(indent);f.left_indent=Cm(0);f.right_indent=Cm(0)
    p.alignment=align
    for r in p._p.iter(qn('w:r')):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        for tag,attrs in [('rFonts',{'ascii':'Book Antiqua','hAnsi':'Book Antiqua','eastAsia':'Book Antiqua','cs':'Book Antiqua'}),('sz',{'val':str(size*2)}),('szCs',{'val':str(size*2)}),('color',{'val':'000000'})]:
            el=pr.find(qn('w:'+tag))
            if el is None:el=OxmlElement('w:'+tag);pr.append(el)
            for k,v in attrs.items():el.set(qn('w:'+k),v)
        if bold is not None:
            el=pr.find(qn('w:b'))
            if el is None:el=OxmlElement('w:b');pr.append(el)
            el.set(qn('w:val'),'1' if bold else '0')

def _remove_body_children(part):
    for child in list(part._element):part._element.remove(child)

def _copy_story(source, target):
    """Retain template anchors and remap image relationships into this package."""
    _remove_body_children(target)
    for child in source._element:
        clone=deepcopy(child)
        for node in clone.iter():
            for attr in (qn('r:embed'),qn('r:id'),qn('r:link')):
                rid=node.get(attr)
                if not rid:continue
                rel=source.part.rels[rid]
                if rel.reltype.endswith('/image') and not rel.is_external:
                    new,_=target.part.get_or_add_image(BytesIO(rel.target_part.blob))
                elif rel.is_external:
                    new=target.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
                else:raise ValueError('Relasi template tidak didukung: '+rel.reltype)
                node.set(attr,new)
        target._element.append(clone)

def _copy_fonts(tpl,doc):
    src=next(r.target_part for r in tpl.part.rels.values() if r.reltype.endswith('/fontTable'))
    dst=next(r.target_part for r in doc.part.rels.values() if r.reltype.endswith('/fontTable'))
    root=etree.fromstring(dst.blob)
    for font in etree.fromstring(src.blob):
        if font.get(qn('w:name')) not in ('Book Antiqua','Garamond'):continue
        for old in list(root):
            if old.get(qn('w:name'))==font.get(qn('w:name')):root.remove(old)
        clone=deepcopy(font)
        for node in clone.iter():
            rid=node.get(qn('r:id'))
            if rid:
                rel=src.rels[rid];part=rel.target_part
                suffix=str(part.partname).rsplit('.',1)[-1]
                name=doc.part.package.next_partname('/word/fonts/rabbani%d.'+suffix)
                copied=Part(name,part.content_type,part.blob,doc.part.package)
                node.set(qn('r:id'),dst.relate_to(copied,rel.reltype))
        root.append(clone)
    dst._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)

def _replace_text(root, replacements):
    # Replace across split runs, without destroying drawings or field codes.
    for p in root.iter(qn('w:p')):
        nodes=p.findall('.//'+qn('w:t'))
        for old,new in replacements.items():
            value=''.join(n.text or '' for n in nodes)
            at=value.find(old)
            if at<0:continue
            end=at+len(old);offset=0;inserted=False
            for n in nodes:
                text=n.text or '';stop=offset+len(text)
                if stop>at and offset<end:
                    prefix=text[:max(0,at-offset)]
                    suffix=text[max(0,end-offset):]
                    n.text=prefix+(new if not inserted else '')+suffix
                    inserted=True
                offset=stop

def _field(p, name):
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),name)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r);p._p.append(field)

def _borders(cell, color=None, fill=None):
    pr=cell._tc.get_or_add_tcPr()
    for tag in ('tcBorders','tcMar','shd'):
        for old in pr.findall(qn('w:'+tag)):pr.remove(old)
    borders=OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if color else 'nil')
        e.set(qn('w:sz'),'6');e.set(qn('w:color'),color or 'FFFFFF');borders.append(e)
    pr.append(borders)
    if fill:
        e=OxmlElement('w:shd');e.set(qn('w:fill'),fill);pr.append(e)
    margins=OxmlElement('w:tcMar')
    for side in ('top','left','bottom','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'100' if color else '0');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)

def _first_footer(part,author,meta):
    _remove_body_children(part)
    table=part.add_table(rows=1,cols=2,width=Cm(15));table.autofit=False
    for c,w in zip(table.columns,(14.3,.7)):c.width=Cm(w)
    left,right=table.rows[0].cells
    left.width=Cm(14.3);right.width=Cm(.7)
    for c in (left,right):_borders(c)
    lines=[f'© ({meta["year"]}) {author}',
           f'Received: {meta["received"]} | Revised: {meta["revised"]} | Accepted: {meta["accepted"]}',
           'Pages: '+meta['pages']]
    for i,line in enumerate(lines):
        p=left.paragraphs[0] if i==0 else left.add_paragraph()
        p.add_run(line);style(p,9,False,1,Align.RIGHT)
    p=right.paragraphs[0];p.add_run('| ');_field(p,'PAGE');style(p,10,False,1,Align.RIGHT)
    p=part.add_paragraph();style(p,1,False,1)

def format_rabbani(source,template,profile='template2023',title_paragraphs=1,metadata=None):
    doc=Document(BytesIO(source));tpl=Document(template);warnings=[]
    _copy_fonts(tpl,doc)
    paragraphs=[p for p in doc.paragraphs if p.text.strip()]
    if len(paragraphs)<5:raise ValueError('Dokumen terlalu pendek. Unggah artikel lengkap, bukan skripsi atau file template.')
    heads=[(p,heading(p.text)) for p in paragraphs if heading(p.text)]
    if not any(h=='intro' for p,h in heads):raise ValueError('Judul bagian PENDAHULUAN / INTRODUCTION tidak ditemukan. Gunakan judul bagian pada paragraf tersendiri.')
    if title_paragraphs not in (1,2,3):raise ValueError('Jumlah paragraf judul harus 1–3.')
    if any(heading(p.text) for p in paragraphs[:title_paragraphs]):raise ValueError('Bagian judul tidak valid.')
    ref=tpl.sections[0]
    article_title=' '.join(p.text.strip() for p in paragraphs[:title_paragraphs])
    author_line=paragraphs[title_paragraphs].text.strip() if len(paragraphs)>title_paragraphs else '[Nama Penulis]'
    meta={k:'—' for k in ('doi','volume','issue','month','year','received','revised','accepted','pages','phone')}
    meta.update({k:str(v).strip() for k,v in (metadata or {}).items() if str(v).strip()})
    running=meta.get('running_title',article_title)
    if len(running)>72:
        running=running[:69].rsplit(' ',1)[0]+'…'
        warnings.append('Judul header diringkas otomatis; isi kolom judul pendek untuk menggantinya.')
    replacements={'Rahmat Dandi, Pendidikan Karakter dalam Islam':author_line+', '+running,
       'DOI: https://doi.org/10.19105/rjpai.v1i1.3002':'DOI: '+meta['doi'],
       'Volume 1, Nomor 1, Maret 2020':f'Volume {meta["volume"]}, Nomor {meta["issue"]}, {meta["month"]} {meta["year"]}',
       'Vol. 1, No. 1, Maret 2020':f'Vol. {meta["volume"]}, No. {meta["issue"]}, {meta["month"]} {meta["year"]}'}
    settings=doc.settings._element
    even_setting=settings.find(qn('w:evenAndOddHeaders'))
    if even_setting is None:
        even_setting=OxmlElement('w:evenAndOddHeaders');settings.append(even_setting)
    even_setting.set(qn('w:val'),'true')
    for sec_index,sec in enumerate(doc.sections):
        for a in ('page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance'):setattr(sec,a,getattr(ref,a))
        cols=sec._sectPr.find(qn('w:cols'))
        if cols is not None:
            cols.set(qn('w:num'),'1')
            for child in list(cols):cols.remove(child)
        sec.different_first_page_header_footer=sec_index==0
        for part in [sec.header,sec.first_page_header,sec.even_page_header,sec.footer,sec.first_page_footer,sec.even_page_footer]:part.is_linked_to_previous=False
        for name in ('first_page_header','header','even_page_header','footer','even_page_footer'):
            part=getattr(sec,name);_copy_story(getattr(ref,name),part)
            _replace_text(part._element,replacements)
        # Keep long running titles clear of the opposite logo.
        for part,alignment in ((sec.header,Align.RIGHT),(sec.even_page_header,Align.LEFT)):
            for p in part.paragraphs:
                p.paragraph_format.left_indent=Cm(2.2 if alignment==Align.RIGHT else 0)
                p.paragraph_format.right_indent=Cm(2.2 if alignment==Align.LEFT else 0)
        _first_footer(sec.first_page_footer,author_line,meta)
    normal=doc.styles['Normal'];normal.font.name='Book Antiqua';normal.font.size=Pt(12)
    state=None;seen=set();abstracts={'abstract':[],'abstrak':[]};keywords=[];front=True
    for i,p in enumerate(paragraphs):
        h=heading(p.text)
        if i<title_paragraphs:
            style(p,13 if profile=='template2023' else 12,True,1.5 if profile=='template2023' else 1,Align.CENTER)
            p.paragraph_format.keep_with_next=True
            if i==0:p.paragraph_format.space_before=Pt(12)
            continue
        if h:
            normalized=re.sub(r'^\s*(?:[A-Z]|\d+|[IVX]+)[.)]\s*','',p.text.strip(),flags=re.I)
            _replace_text(p._p,{p.text:normalized})
            state=h;seen.add(h);front=False
            style(p,11 if h in abstracts else 12,True,1 if h in abstracts or h=='references' else 1.5,Align.LEFT)
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
        elif re.match(r'^(keywords?|kata\s*kunci)\s*:',p.text,re.I):
            keywords.append(p);style(p,11,spacing=1,align=Align.LEFT)
        elif front:
            small=bool(re.search(r'@|universitas|institut|sekolah tinggi|university|\bUIN\b|\bIAIN\b',p.text,re.I))
            style(p,10 if small else 11,True,1.5,Align.CENTER)
            for r in p.runs:r.italic=False
        elif state in abstracts:
            abstracts[state].append(p);style(p,11,spacing=1)
            for r in p.runs:r.italic=False
        elif state=='references':
            category=bool(re.match(r'^(buku|book|artikel\s+dalam\s+jurnal|artikel\s+jurnal|journal\s+articles?)\s*$',p.text.strip(),re.I))
            if category:
                style(p,12,True,1,Align.LEFT);p.paragraph_format.left_indent=Cm(0);p.paragraph_format.first_line_indent=Cm(0)
                for r in p.runs:
                    r.underline=True
            else:
                style(p,12,False,1,Align.LEFT);p.paragraph_format.left_indent=Cm(1);p.paragraph_format.first_line_indent=Cm(-1);p.paragraph_format.space_after=Pt(5)
        else:
            style(p,12,spacing=1.5,indent=1.5)
    title=' '.join(p.text for p in paragraphs[:title_paragraphs])
    if len(title.split())>20:warnings.append('Judul melebihi 20 kata; isi judul dipertahankan.')
    for h,ps in abstracts.items():
        words=len(' '.join(p.text for p in ps).split())
        if not ps:warnings.append(f'{h.upper()} belum ditemukan; tidak diterjemahkan atau dibuat otomatis.')
        elif not ((150<=words<=200) if profile=='template2023' else (130<=words<=170)):warnings.append(f'{h.upper()}: {words} kata; periksa panjang sesuai panduan.')
        if len(ps)>1:warnings.append(f'{h.upper()} terdiri dari {len(ps)} paragraf; panduan meminta satu paragraf.')
    for h in ['method','results','conclusion','references']:
        if h not in seen:warnings.append('Bagian belum terdeteksi: '+h)
    if not keywords:warnings.append('Kata kunci belum ditemukan. Gunakan label Keywords: / Kata kunci:.')
    for p in keywords:
        values=re.split(r'[,;]',p.text.split(':',1)[1]);n=len([x for x in values if x.strip()])
        if not 3<=n<=6:warnings.append(f'Kata kunci terdeteksi {n} butir; gunakan 3–6 dipisahkan koma/titik koma.')
    # Reproduce the reference abstract/sidebar using flow layout, so long text can paginate.
    abstract_nodes=[p for p in paragraphs if heading(p.text) in abstracts or any(p is x for ps in abstracts.values() for x in ps)]
    if profile=='template2023' and abstract_nodes:
        # A single splittable row avoids the merged-cell pagination failure.
        # The separate nested sidebar has its own bottom border and height.
        box=doc.add_table(rows=1,cols=3);box.autofit=False
        split=OxmlElement('w:cantSplit');split.set(qn('w:val'),'0');box.rows[0]._tr.get_or_add_trPr().append(split)
        for col,cell,width in zip(box.columns,box.rows[0].cells,(9.4,.16,5.44)):
            col.width=Cm(width);cell.width=Cm(width);_borders(cell)
        abstract_nodes[0]._p.addprevious(box._tbl)
        left=box.cell(0,0)
        left._tc.remove(left.paragraphs[0]._p)
        for h in ('abstract','abstrak'):
            group=[p for p in abstract_nodes if heading(p.text)==h]+abstracts[h]
            for p in group:
                left._tc.append(p._p)
                if heading(p.text):
                    p.paragraph_format.space_before=Pt(11 if h=='abstrak' else 0)
                    p.paragraph_format.keep_with_next=False
        _borders(left,'45A9C4','D4F6FF')
        sidebar=box.cell(0,2).add_table(rows=1,cols=1);sidebar.autofit=False
        box.cell(0,2)._tc.remove(box.cell(0,2).paragraphs[0]._p)
        sidebar.columns[0].width=Cm(5.44);right=sidebar.cell(0,0);right.width=Cm(5.44)
        _borders(right,'00B0F0','FFFFFF')
        p=right.paragraphs[0];p.add_run('ARTICLE HISTORY');style(p,11,True,1,Align.LEFT)
        for label,key in [('Received','received'),('Revised','revised'),('Accepted','accepted')]:
            p=right.add_paragraph(label+' '+meta[key]);style(p,11,False,1,Align.LEFT)
        p=right.add_paragraph('KEYWORDS');style(p,11,True,1,Align.LEFT);p.paragraph_format.space_before=Pt(11)
        for p in sorted(keywords,key=lambda p:not p.text.lower().startswith('keyword')):
            right._tc.append(p._p);style(p,11,False,1,Align.LEFT)
            for r in p.runs:r.italic=False
        for c in (box.cell(0,1),box.cell(0,2)):
            for p in c.paragraphs:
                style(p,1,False,1);p.paragraph_format.line_spacing=Pt(1)
        # Preserve the contact slot even when the source lacks a phone number.
        front_text=' '.join(p.text for p in paragraphs[:next(i for i,p in enumerate(paragraphs) if heading(p.text))])
        if not re.search(r'\b(?:WA|WhatsApp|08\d{7,})\b',front_text,re.I):
            p=doc.add_paragraph('Nomor WA yang bisa dihubungi: '+meta['phone'])
            style(p,12,True,1.5,Align.CENTER);p.paragraph_format.space_after=Pt(6)
            box._tbl.addprevious(p._p)
    warnings.append('Metadata bertanda — belum tersedia. Lengkapi DOI, edisi, tanggal editorial, halaman terbit dan nomor WA sebelum pengiriman.')
    warnings.append('Sitasi Chicago 17 Full Footnote dan isi ilmiah perlu diperiksa; referensi tidak direkayasa atau dikonversi otomatis.')
    warnings.append('Tabel/gambar dan footnote asli dipertahankan; periksa lebar tabel/gambar dan pemenggalan halaman di Word.')
    doc.core_properties.comments='RABBANI: periksa catatan validasi pada halaman web. Sitasi dan isi sumber dipertahankan.'
    return doc,warnings
