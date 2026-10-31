"""Al-Ihkam 2026: layout measured from the user's original seven-page template."""
from io import BytesIO
from copy import deepcopy
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .rabbani import heading, style, _copy_story, _copy_fonts, _replace_text

EXTRA = {'declaration of generative ai and ai-assisted technologies':'ai',
         'acknowledgment':'ack','acknowledgments':'ack','acknowledgement':'ack',
         'ucapan terima kasih':'ack','deklarasi penggunaan ai':'ai',
         'result and discussion':'results','method':'method'}

def section_name(text):
    clean=re.sub(r'\s*\([^)]*paragraphs[^)]*\)\s*$', '', text.strip(),flags=re.I)
    return EXTRA.get(clean.lower().rstrip(':')) or heading(clean)

def _format(p, sample, size=11, bold=None, italic=None):
    style(p,size,bold,1)
    old=p._p.pPr
    section=deepcopy(old.find(qn('w:sectPr'))) if old is not None and old.find(qn('w:sectPr')) is not None else None
    if old is not None:p._p.remove(old)
    if sample._p.pPr is not None:p._p.insert(0,deepcopy(sample._p.pPr))
    if section is not None:p._p.get_or_add_pPr().append(section)
    f=p.paragraph_format
    f.line_spacing=1;f.space_before=Pt(0);f.space_after=Pt(0)
    f.keep_with_next=False;f.keep_together=False;f.page_break_before=False
    if italic is not None:
        for r in p.runs:r.italic=italic

def format_alihkam(source,template,title_paragraphs=1,metadata=None):
    doc=Document(BytesIO(source));tpl=Document(template);warnings=[]
    ps=[p for p in doc.paragraphs if p.text.strip()]
    if title_paragraphs not in (1,2,3) or len(ps)<5:
        raise ValueError('Artikel harus memiliki judul, penulis, dan isi; jumlah paragraf judul 1–3.')
    if not any(section_name(p.text)=='intro' for p in ps):
        raise ValueError('Bagian Introduction / PENDAHULUAN tidak ditemukan.')
    meta={k:'—' for k in ('volume','issue','year','pages','doi','received','revised','accepted','published')}
    meta.update({k:str(v).strip() for k,v in (metadata or {}).items() if str(v).strip()})
    title=' '.join(p.text.strip() for p in ps[:title_paragraphs])
    author=ps[title_paragraphs].text.strip()
    front_end=next(i for i,p in enumerate(ps) if section_name(p.text))
    front=' '.join(p.text for p in ps[:front_end])
    email_match=re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}',front)
    email=meta.get('correspondence_email',email_match.group() if email_match else '—')
    short=meta.get('running_title',title)
    if len(short)>65:
        short=short[:62].rsplit(' ',1)[0]+'…'
        warnings.append('Judul header dipendekkan; dapat diganti melalui kolom judul pendek.')
    edition=f'{meta["volume"]} ({meta["issue"]}), {meta["year"]}: {meta["pages"]}'
    replacements={'The Title Article Journal …………………….':short,
        'full name of first author, et al.':meta.get('running_author',author),
        '21 (1), 2026: 1-25':edition,
        'https://doi.org/10.19105/al-lhkam.v17i2.7229':meta['doi'],
        'address@mail.ac.id':email,
        'Copyright (c) 2026':'Copyright (c) '+meta['year']}
    _copy_fonts(tpl,doc)
    normal=doc.styles['Normal'];normal.font.name='Book Antiqua';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1
    normal.paragraph_format.space_after=Pt(0)
    setting=doc.settings._element.find(qn('w:evenAndOddHeaders'))
    if setting is None:setting=OxmlElement('w:evenAndOddHeaders');doc.settings._element.append(setting)
    setting.set(qn('w:val'),'true')
    ref=tpl.sections[0]
    for i,sec in enumerate(doc.sections):
        for name in ('page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance'):
            setattr(sec,name,getattr(ref,name))
        sec.different_first_page_header_footer=i==0
        cols=sec._sectPr.find(qn('w:cols'))
        if cols is not None:
            cols.set(qn('w:num'),'1')
            for child in list(cols):cols.remove(child)
        for name in ('header','even_page_header','first_page_header','footer','even_page_footer','first_page_footer'):
            dst=getattr(sec,name);dst.is_linked_to_previous=False
            _copy_story(getattr(ref,name),dst);_replace_text(dst._element,replacements)
            # Update clickable targets as well as visible text, including email.
            for rel in dst.part.rels.values():
                if rel.is_external:
                    if '10.19105/al-lhkam.v17i2.7229' in rel.target_ref:
                        rel._target=meta['doi'] if meta['doi'].startswith('https://') else ('https://doi.org/'+meta['doi'] if meta['doi']!='—' else 'https://ejournal.uinmadura.ac.id/index.php/alihkam/')
                    elif 'address@mail.ac.id' in rel.target_ref:
                        rel._target='mailto:'+email if email!='—' else 'https://ejournal.uinmadura.ac.id/index.php/alihkam/'
    state=None;counts={};abstract=[];seen=set();keyword_count=0;history_existing=None
    samples=tpl.paragraphs
    for i,p in enumerate(ps):
        h=section_name(p.text)
        if i<title_paragraphs:
            _format(p,samples[3],13,True,False);p.alignment=Align.CENTER
            p.paragraph_format.keep_with_next=True
            if i==0:p.paragraph_format.space_before=Pt(13)
            if i==title_paragraphs-1:p.paragraph_format.space_after=Pt(11)
        elif h:
            state=h;seen.add(h)
            _format(p,samples[18] if h in ('abstract','abstrak') else samples[27],10 if h in ('abstract','abstrak') else 11,True,False)
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(11)
        elif re.match(r'^(keywords?|kata\s*kunci)\s*:',p.text,re.I) or state=='keywords':
            value=p.text.split(':',1)[1] if re.match(r'^(keywords?|kata\s*kunci)\s*:',p.text,re.I) else p.text
            values=[v.strip() for v in re.split('[;,]',value) if v.strip()]
            if state!='keywords' and values:
                label=p.insert_paragraph_before('Keywords:');_format(label,samples[22],10,True,False)
                label.paragraph_format.space_before=Pt(10);label.paragraph_format.keep_with_next=True
            if not values:
                _format(p,samples[22],10,True,False);state='keywords';continue
            keyword_count=len(values)
            _replace_text(p._p,{p.text:'; '.join(sorted(values,key=str.casefold))})
            _format(p,samples[24],10,False,False);p.alignment=Align.CENTER
            p.paragraph_format.space_after=Pt(11);state='after_keywords'
        elif re.match(r'^Article history\s*:',p.text,re.I):
            history_existing=p
        elif state is None:
            affiliation=bool(re.search(r'@|universitas|institut|university|\bUIN\b|\bIAIN\b|program studi',p.text,re.I))
            _format(p,samples[7] if affiliation else samples[6],10 if affiliation else 11,not affiliation,affiliation)
        elif state in ('abstract','abstrak'):
            _format(p,samples[20],10,False);abstract.append((state,p))
        elif state=='references':
            _format(p,samples[72],11,False)
        elif p.style.name.lower().startswith('heading'):
            _format(p,samples[38],11,True,True);p.paragraph_format.keep_with_next=True
        elif re.match(r'^(Table|Tabel|Figure|Gambar)\s+\d+[.:]',p.text,re.I):
            table=bool(re.match(r'^(Table|Tabel)',p.text,re.I))
            _format(p,samples[43] if table else samples[49],11)
            p.paragraph_format.keep_with_next=table
        else:
            quote=p.text.lstrip().startswith(('“','"')) and p.text.rstrip().endswith(('”','"',')'))
            _format(p,samples[57] if quote else samples[28],11)
            counts[state]=counts.get(state,0)+1
    first=ps[front_end]
    history=history_existing or first.insert_paragraph_before(f'Article history: Received: {meta["received"]}; Revised: {meta["revised"]}; Accepted: {meta["accepted"]}; Published: {meta["published"]}')
    _format(history,samples[15],8,False,False);history.alignment=Align.CENTER
    history.paragraph_format.space_before=Pt(11)
    for lang,p in abstract:
        if len(p.text.split())>250:warnings.append(f'{lang.upper()} melebihi 250 kata; teks tidak dipotong.')
    if 'abstract' not in seen:warnings.append('Abstract bahasa Inggris belum ditemukan; tidak diterjemahkan otomatis.')
    if 'abstrak' in seen:warnings.append('ABSTRAK Indonesia ikut dipertahankan; template Al-Ihkam mencontohkan Abstract bahasa Inggris saja.')
    if not 3<=keyword_count<=5:warnings.append('Template meminta 3–5 kata kunci, urut alfabet, dipisahkan titik koma.')
    for h,label,lo,hi in [('intro','Introduction',4,6),('method','Methods',3,3),('results','Result and Discussion',20,25),('conclusion','Conclusion',1,2)]:
        if not lo<=counts.get(h,0)<=hi:warnings.append(f'{label}: {counts.get(h,0)} paragraf terdeteksi; template meminta {lo}–{hi}.')
    for h,label in [('ai','Declaration of Generative AI and AI-assisted Technologies'),('ack','Acknowledgment'),('references','Bibliography')]:
        if h not in seen:warnings.append(f'Bagian {label} belum ada; isi pernyataan tidak dibuat otomatis.')
    warnings.append('Metadata — perlu dilengkapi. Sitasi Chicago 17 dan DOI referensi perlu diperiksa; isi ilmiah dan footnote sumber dipertahankan.')
    return doc,warnings
