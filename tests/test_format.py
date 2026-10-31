"""Run: python -m unittest discover -s tests (requires httpx)."""
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from backend.main import app, ROOT

class FormatTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app)
        d=Document()
        for t in ['PENDIDIKAN KARAKTER DI SEKOLAH','Ahmad Zain','UIN Madura','penulis@example.com','ABSTRACT','This is a formatting test.','ABSTRAK','Ini adalah uji pemformatan.','Keywords: character; education; learning','A. PENDAHULUAN','Isi pendahuluan tetap utuh.','B. METODE PENELITIAN','Metode contoh.','C. HASIL DAN PEMBAHASAN','Hasil contoh.','D. KESIMPULAN','Kesimpulan contoh.','DAFTAR PUSTAKA','Penulis. Buku. Kota: Penerbit, 2023.']:d.add_paragraph(t)
        d.add_table(rows=1,cols=1).cell(0,0).text='Tabel sumber'
        b=BytesIO();d.save(b);self.source=b.getvalue()
    def post(self,tid,profile='template2023',source=None):
        return self.client.post('/format',files={'file':('artikel.docx',source or self.source)},data={'template_id':tid,'profile':profile})
    def test_templates_and_profiles(self):
        self.assertEqual({t['id'] for t in self.client.get('/templates').json()},{'karsa','rabbani','alihkam'})
        for tid,profile in [('karsa','template2023'),('rabbani','template2023'),('rabbani','guidelines')]:
            with self.subTest(tid=tid,profile=profile):
                r=self.post(tid,profile);self.assertEqual(r.status_code,200,r.text[:300] if r.status_code!=200 else '')
                d=Document(BytesIO(r.content));text=' '.join(p.text for p in d.paragraphs)
                self.assertIn('Isi pendahuluan tetap utuh.',text)
                self.assertIn('Tabel sumber',' '.join(c.text for t in d.tables for row in t.rows for c in row.cells))
                self.assertIn(tid.upper(),r.headers['content-disposition'])
                if tid=='rabbani':
                    self.assertAlmostEqual(d.sections[0].left_margin.cm,3,places=2)
                    self.assertEqual(d.paragraphs[0].runs[0].font.size.pt,13 if profile=='template2023' else 12)
                    self.assertNotIn('Karsa',d.sections[0].first_page_header._element.xml)
    def test_original_footnotes_and_media_survive(self):
        source=(ROOT/'templates/rabbani/template.docx').read_bytes()
        r=self.post('rabbani',source=source);self.assertEqual(r.status_code,200)
        with ZipFile(BytesIO(source)) as a,ZipFile(BytesIO(r.content)) as b:
            names=[n for n in a.namelist() if n.startswith('word/media/') or n=='word/footnotes.xml']
            self.assertTrue(names)
            for name in names:self.assertEqual(a.read(name),b.read(name),name)
    def test_reject_invalid_requests(self):
        self.assertEqual(self.post('../karsa').status_code,404)
        self.assertEqual(self.post('rabbani','unknown').status_code,400)
        self.assertEqual(self.post('rabbani',source=b'not a docx').status_code,422)
        r=self.client.post('/format',files={'file':('input.pdf',b'pdf')});self.assertEqual(r.status_code,400)
    def test_rabbani_parity_fonts_and_panel(self):
        r=self.post('rabbani');self.assertEqual(r.status_code,200)
        d=Document(BytesIO(r.content));s=d.sections[0]
        self.assertTrue(s.different_first_page_header_footer)
        self.assertEqual(d.settings._element.find(qn('w:evenAndOddHeaders')).get(qn('w:val')),'true')
        for name,align in [('header','right'),('even_page_header','left'),('first_page_header','left')]:
            story=getattr(s,name)
            self.assertEqual(story.paragraphs[0]._p.pPr.find(qn('w:jc')).get(qn('w:val')),align)
            self.assertTrue(story._element.xpath('.//a:blip'))
            self.assertNotIn('Rahmat Dandi',story._element.xml)
            for blip in story._element.xpath('.//a:blip'):
                self.assertTrue(story.part.rels[blip.get(qn('r:embed'))].target_part.blob)
        even=s.even_page_footer._element.xml
        odd=s.footer._element.xml
        self.assertLess(even.index('PAGE'),even.index('Rabbani'))
        self.assertLess(odd.index('Rabbani'),odd.index('PAGE'))
        self.assertIn('Pages:',s.first_page_footer._element.xml)
        panel=d.tables[0]
        self.assertTrue(panel.cell(0,0).text.startswith('ABSTRACT'))
        self.assertIn('ABSTRAK',panel.cell(0,0).text)
        self.assertEqual(panel.rows[0]._tr.trPr.find(qn('w:cantSplit')).get(qn('w:val')),'0')
        self.assertFalse(panel.cell(0,0).paragraphs[0].paragraph_format.keep_with_next)
        self.assertEqual(len(panel.cell(0,2).tables),1)
        fonts=next(x.target_part for x in d.part.rels.values() if x.reltype.endswith('/fontTable'))
        self.assertGreaterEqual(len(fonts.rels),4)
    def test_metadata_not_taken_from_bibliography(self):
        r=self.client.post('/format',files={'file':('a.docx',self.source)},data={'template_id':'rabbani','metadata':'{"doi":"10.example/article","running_title":"Judul pendek","year":"2026"}'})
        self.assertEqual(r.status_code,200)
        d=Document(BytesIO(r.content))
        self.assertIn('DOI: 10.example/article',d.sections[0].header._element.xml)
        self.assertIn('Ahmad Zain, Judul pendek',d.sections[0].header._element.xml)
        self.assertIn('2026',d.sections[0].first_page_footer._element.xml)
        r=self.client.post('/format',files={'file':('a.docx',self.source)},data={'template_id':'rabbani','metadata':'[]'})
        self.assertEqual(r.status_code,422)
    def test_alihkam_layout_and_metadata(self):
        r=self.post('alihkam','template2026');self.assertEqual(r.status_code,200,r.text[:200] if r.status_code!=200 else '')
        d=Document(BytesIO(r.content));s=d.sections[0]
        self.assertAlmostEqual(s.page_width.cm,21.59,places=2)
        self.assertAlmostEqual(s.top_margin.cm,5,places=2)
        self.assertAlmostEqual(s.right_margin.cm,4.5,places=2)
        self.assertIn('Ahmad Zain',s.even_page_header._element.xml)
        self.assertIn('PENDIDIKAN KARAKTER DI SEKOLAH',s.header._element.xml)
        self.assertNotIn('10.19105/al-lhkam.v17i2.7229',s.first_page_header._element.xml)
        self.assertNotIn('21 (1), 2026: 1-25',s.first_page_header._element.xml)
        self.assertIn('penulis@example.com',s.first_page_footer._element.xml)
        self.assertEqual(d.paragraphs[0].runs[0].font.size.pt,13)
        self.assertIn('Article history:', '\n'.join(p.text for p in d.paragraphs))
        self.assertEqual(len(d.tables),1) # Source table only; no Rabbani panel.
        self.assertEqual(self.post('alihkam','guidelines').status_code,400)
    def test_alihkam_separate_keywords_label(self):
        d=Document(BytesIO(self.source))
        p=next(p for p in d.paragraphs if p.text.startswith('Keywords:'))
        p.insert_paragraph_before('Keywords:');p.text='learning; character; education'
        b=BytesIO();d.save(b)
        r=self.post('alihkam','template2026',b.getvalue());self.assertEqual(r.status_code,200)
        out=Document(BytesIO(r.content))
        p=next(p for p in out.paragraphs if p.text=='character; education; learning')
        self.assertEqual(p.alignment,1)
        self.assertEqual(p.runs[0].font.size.pt,10)
    def test_cors_exposes_validation_notes(self):
        r=self.client.post('/format',files={'file':('article.docx',self.source)},data={'template_id':'rabbani'},headers={'Origin':'https://example.com'})
        self.assertIn('X-Format-Warnings',r.headers['access-control-expose-headers'])

if __name__=='__main__':unittest.main()
